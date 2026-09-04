from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import csv
import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.db.database import get_db
from app.schemas.usuarios import (
    UsuarioRead, UsuarioCreate, UsuarioUpdate,
    PagaRead, PagaCreate, RegistroMensualCreate,
)
from app.services import usuarios as svc
from app.api.contabilidad import _csv_response, _xlsx_response

_MESES_ABR = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

router = APIRouter(prefix="/api/usuarios", tags=["usuarios"])


# ─── Usuarios NNA ─────────────────────────────────────────────────────────────

@router.get("", response_model=dict)
def listar(
    empresa_id: int,
    activo: Optional[bool] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    items, total = svc.get_usuarios(db, empresa_id, activo, skip, limit)
    return {"total": total, "items": [UsuarioRead.model_validate(i) for i in items]}


@router.post("", response_model=UsuarioRead, status_code=201)
def crear(data: UsuarioCreate, db: Session = Depends(get_db)):
    return svc.create_usuario(db, data)


@router.put("/{usuario_id}", response_model=UsuarioRead)
def actualizar(usuario_id: int, data: UsuarioUpdate, db: Session = Depends(get_db)):
    u = svc.update_usuario(db, usuario_id, data)
    if not u:
        raise HTTPException(404, "Usuario no encontrado")
    return u


@router.delete("/{usuario_id}", status_code=204)
def eliminar(usuario_id: int, db: Session = Depends(get_db)):
    try:
        if not svc.delete_usuario(db, usuario_id):
            raise HTTPException(404, "Usuario no encontrado")
    except ValueError as e:
        db.rollback()
        raise HTTPException(400, str(e))


@router.get("/activos-con-paga", response_model=list[UsuarioRead])
def activos_con_paga(empresa_id: int, db: Session = Depends(get_db)):
    return svc.get_activos_con_paga(db, empresa_id)


@router.get("/saldos", response_model=dict)
def saldos_nna(empresa_id: int, db: Session = Depends(get_db)):
    """Devuelve {numero: saldo} desde diario para las cuentas 4001xxx."""
    return svc.get_saldos_nna(db, empresa_id)


# ─── Pagas ────────────────────────────────────────────────────────────────────

@router.get("/pagas", response_model=dict)
def listar_pagas(
    empresa_id: int,
    usuario: Optional[int] = None,
    fecha_desde: Optional[datetime.date] = None,
    fecha_hasta: Optional[datetime.date] = None,
    skip: int = 0,
    limit: int = 100,
    sort_by: str = "fecha",
    sort_dir: str = "desc",
    db: Session = Depends(get_db),
):
    items, total = svc.get_pagas(db, empresa_id, usuario, fecha_desde, fecha_hasta, skip, limit,
                                  sort_by, sort_dir)
    return {"total": total, "items": [PagaRead.model_validate(i) for i in items]}


@router.post("/pagas", response_model=PagaRead, status_code=201)
def crear_paga(data: PagaCreate, db: Session = Depends(get_db)):
    return svc.create_paga(db, data)


@router.delete("/pagas/{paga_id}", status_code=204)
def eliminar_paga(paga_id: int, db: Session = Depends(get_db)):
    try:
        if not svc.delete_paga(db, paga_id):
            raise HTTPException(404, "Paga no encontrada")
    except ValueError as e:
        db.rollback()
        raise HTTPException(400, str(e))


@router.post("/pagas/mes", response_model=list[PagaRead], status_code=201)
def registrar_mes(data: RegistroMensualCreate, db: Session = Depends(get_db)):
    return svc.registrar_mes(db, data)


@router.get("/pagas/export")
def exportar_pagas(
    empresa_id: int,
    usuario: Optional[int] = None,
    fecha_desde: Optional[datetime.date] = None,
    fecha_hasta: Optional[datetime.date] = None,
    format: str = "xlsx",
    sort_by: str = "fecha",
    sort_dir: str = "desc",
    db: Session = Depends(get_db),
):
    """Exporta el listado de pagas a Excel (.xlsx) o CSV (.csv)."""
    items, _ = svc.get_pagas(db, empresa_id, usuario, fecha_desde, fecha_hasta, skip=0, limit=10000,
                              sort_by=sort_by, sort_dir=sort_dir)
    usuarios_map = {u.numero: f"{u.nombre} {u.apellidos or ''}".strip()
                    for u in svc.get_usuarios(db, empresa_id, limit=1000)[0]}

    nombre_base = f"pagas_nna_{empresa_id}"
    if fecha_desde:
        nombre_base += f"_{fecha_desde}"
    if fecha_hasta:
        nombre_base += f"_{fecha_hasta}"

    filas = [
        (p.fecha.strftime('%d/%m/%Y') if p.fecha else '',
         usuarios_map.get(p.usuario, f"NNA {p.usuario}"), float(p.importe))
        for p in items
    ]
    total = sum(f[2] for f in filas)

    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=";")
        writer.writerow(["Fecha", "Usuario", "Importe"])
        for fecha, nombre, importe in filas:
            writer.writerow([fecha, nombre, f"{importe:.2f}".replace(".", ",")])
        writer.writerow(["TOTAL", "", f"{total:.2f}".replace(".", ",")])
        return StreamingResponse(
            iter([buf.getvalue().encode("utf-8-sig")]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{nombre_base}.csv"'},
        )

    # Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pagas NNA"

    hdr_fill = PatternFill("solid", fgColor="1a5276")
    hdr_font = Font(bold=True, color="FFFFFF")
    for col, texto in enumerate(["Fecha", "Usuario", "Importe"], start=1):
        c = ws.cell(row=1, column=col, value=texto)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 14

    for row_idx, (fecha, nombre, importe) in enumerate(filas, start=2):
        ws.cell(row=row_idx, column=1, value=fecha)
        ws.cell(row=row_idx, column=2, value=nombre)
        c = ws.cell(row=row_idx, column=3, value=importe)
        c.number_format = '#,##0.00 €'

    if filas:
        r = len(filas) + 2
        ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=total)
        c.number_format = '#,##0.00 €'
        c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_base}.xlsx"'},
    )


@router.get("/pagas/resumen", response_model=dict)
def resumen_pagas(
    empresa_id: int,
    fecha_desde: datetime.date,
    fecha_hasta: datetime.date,
    db: Session = Depends(get_db),
):
    return svc.resumen_pagas(db, empresa_id, fecha_desde, fecha_hasta)


@router.get("/pagas/anios", response_model=dict)
def anios_pagas(empresa_id: int, db: Session = Depends(get_db)):
    return {"anios": svc.anios_pagas(db, empresa_id)}


@router.get("/pagas/resumen-anual", response_model=dict)
def resumen_pagas_anual(empresa_id: int, anio: int, db: Session = Depends(get_db)):
    return svc.resumen_pagas_anual(db, empresa_id, anio)


@router.get("/pagas/resumen-anual/export")
def exportar_resumen_pagas_anual(
    empresa_id: int, anio: int, format: str = 'xlsx', db: Session = Depends(get_db),
):
    data = svc.resumen_pagas_anual(db, empresa_id, anio)
    rows = [['NNA'] + _MESES_ABR + ['Total']]
    for u in data['usuarios']:
        rows.append([u['nombre']] + u['meses'] + [u['total']])
    rows.append(['Total mes'] + data['totales_mes'] + [data['total_anual']])
    nombre = f"pagas_nna_{anio}.{format}"
    return _xlsx_response(rows, nombre) if format == 'xlsx' else _csv_response(rows, nombre)
